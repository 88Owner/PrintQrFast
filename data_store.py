"""Master data (SKU, product name, QR payload) and print settings."""

from __future__ import annotations

import copy
import csv
import io
import json
import threading
import unicodedata
from pathlib import Path

LABEL_WIDTH_MM = 35.0
LABEL_HEIGHT_MM = 22.0
LABEL_GAP_MM = 2.0
ROW_GAP_MM = 3.0
SHEET_WIDTH_MM = 74.0
SHEET_HEIGHT_MM = 22.0
FORM_HEIGHT_MM = SHEET_HEIGHT_MM + ROW_GAP_MM

DEFAULT_SETTINGS = {
    "printer": "",
    "preset": "35x22",
    "label_width_mm": LABEL_WIDTH_MM,
    "label_height_mm": LABEL_HEIGHT_MM,
    "copies": 1,
}

SKU_HEADERS = {"sku", "ma sku", "ma san pham", "masanpham", "ma hang"}
NAME_HEADERS = {"ten sp", "ten san pham", "ten", "name", "product", "product name", "san pham", "ten hang"}
QR_HEADERS = {"qr code", "qr", "ma qr", "qrcode", "noi dung qr", "noi dung ma qr"}


def normalize_copies(copies) -> int:
    """Sticker count. A row is two labels, so the count is always even."""
    try:
        count = int(copies)
    except (TypeError, ValueError):
        count = 2
    if count < 2:
        count = 2
    if count % 2:
        count += 1
    if count > 100:
        count = 100
    if count % 2:
        count -= 1
    return count


def fmt_num(value: float) -> str:
    number = float(value)
    if abs(number - round(number)) < 1e-9:
        return str(int(round(number)))
    return f"{number:.2f}".rstrip("0").rstrip(".")


def fold(value: str) -> str:
    text = unicodedata.normalize("NFD", str(value))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = text.replace("_", " ").replace("-", " ")
    return " ".join(text.casefold().split())


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def clean_product(body: dict) -> dict:
    sku = " ".join(cell_text(body.get("sku", "")).split())
    name = " ".join(cell_text(body.get("name", "")).split())
    qr = cell_text(body.get("qr", "")).strip()
    if not sku:
        raise ValueError("Thiếu SKU")
    if len(sku) > 80:
        raise ValueError("SKU tối đa 80 ký tự")
    if not name:
        raise ValueError("Thiếu tên sản phẩm")
    if len(name) > 160:
        raise ValueError("Tên sản phẩm tối đa 160 ký tự")
    if len(qr) > 1000:
        raise ValueError("QR code tối đa 1000 ký tự")
    if not qr:
        qr = sku
    return {"sku": sku, "name": name, "qr": qr}


def parse_label(body: dict) -> tuple[dict, float, float]:
    return clean_product(body), LABEL_WIDTH_MM, LABEL_HEIGHT_MM


def default_master() -> dict:
    return {"products": []}


def master_view(data: dict) -> dict:
    products = sorted(data["products"], key=lambda product: (product["name"].casefold(), product["sku"].casefold()))
    return {"products": products}


def parse_table(filename: str, raw: bytes) -> list[dict]:
    if not raw:
        raise ValueError("File trống")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        rows = _read_xlsx(raw)
    elif name.endswith(".xls"):
        raise ValueError("Hãy lưu file Excel dạng .xlsx hoặc CSV")
    else:
        rows = _read_delimited(raw)
    return _rows_to_products(rows)


def _read_xlsx(raw: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Chưa cài openpyxl. Chạy lại start.bat.") from exc
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Không đọc được file Excel") from exc
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([cell_text(value) for value in row])
    workbook.close()
    return rows


def _read_delimited(raw: bytes) -> list[list[str]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1258"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    sample = text[:2000]
    delimiter = "\t" if sample.count("\t") >= sample.count(",") and sample.count("\t") >= sample.count(";") else None
    if delimiter is None:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell.strip() for cell in row] for row in reader]


def _header_index(header: list[str], names: set[str]) -> int | None:
    for index, cell in enumerate(header):
        if fold(cell) in names:
            return index
    return None


def _rows_to_products(rows: list[list[str]]) -> list[dict]:
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError("File không có dữ liệu")
    header = [fold(cell) for cell in rows[0]]
    has_header = any(cell in SKU_HEADERS or cell in NAME_HEADERS or cell in QR_HEADERS for cell in header)
    if has_header:
        sku_index = _header_index(header, SKU_HEADERS)
        name_index = _header_index(header, NAME_HEADERS)
        qr_index = _header_index(header, QR_HEADERS)
        if sku_index is None or name_index is None:
            raise ValueError("File cần có cột SKU và Tên SP")
        data_rows = rows[1:]
        first_number = 2
    else:
        if len(rows[0]) < 2:
            raise ValueError("Mỗi dòng cần SKU và Tên SP")
        sku_index, name_index, qr_index = 0, 1, 2 if len(rows[0]) > 2 else None
        data_rows = rows
        first_number = 1
    products = []
    seen: dict[str, int] = {}
    for offset, row in enumerate(data_rows):
        line_no = first_number + offset
        sku = row[sku_index] if sku_index < len(row) else ""
        name = row[name_index] if name_index < len(row) else ""
        qr = row[qr_index] if qr_index is not None and qr_index < len(row) else ""
        if not sku.strip() and not name.strip() and not str(qr).strip():
            continue
        try:
            product = clean_product({"sku": sku, "name": name, "qr": qr})
        except ValueError as exc:
            raise ValueError(f"Dòng {line_no}: {exc}") from exc
        key = product["sku"].casefold()
        if key in seen:
            continue
        seen[key] = line_no
        products.append(product)
    if not products:
        raise ValueError("File không có sản phẩm")
    return products


class Store:
    def __init__(self, root: Path):
        self.data_dir = root / "data"
        self.master_path = self.data_dir / "master.json"
        self.settings_path = self.data_dir / "settings.json"
        self.lock = threading.RLock()

    def load_master(self) -> dict:
        with self.lock:
            return master_view(self._read_master())

    def lookup(self, code: str) -> dict | None:
        key = str(code).strip()
        if not key:
            return None
        folded = key.casefold()
        with self.lock:
            for product in self._read_master()["products"]:
                if product["sku"].casefold() == folded or product["qr"].casefold() == folded:
                    return copy.deepcopy(product)
        return None

    def upsert_product(self, body: dict) -> dict:
        product = clean_product(body)

        def mutate(data: dict) -> dict:
            for index, current in enumerate(data["products"]):
                if current["sku"].casefold() == product["sku"].casefold():
                    data["products"][index] = product
                    return product
            data["products"].append(product)
            return product

        return self._update_master(mutate)

    def delete_product(self, sku: str) -> None:
        key = str(sku).strip().casefold()

        def mutate(data: dict) -> None:
            before = len(data["products"])
            data["products"] = [product for product in data["products"] if product["sku"].casefold() != key]
            if len(data["products"]) == before:
                raise ValueError("Không tìm thấy SKU")

        self._update_master(mutate)

    def import_products(self, products: list[dict], replace: bool = True) -> dict:
        cleaned = []
        seen: set[str] = set()
        for product in products:
            item = clean_product(product)
            key = item["sku"].casefold()
            if key in seen:
                continue
            seen.add(key)
            cleaned.append(item)

        def mutate(data: dict) -> dict:
            if replace:
                data["products"] = cleaned
            else:
                index_by_sku = {product["sku"].casefold(): index for index, product in enumerate(data["products"])}
                for product in cleaned:
                    key = product["sku"].casefold()
                    if key in index_by_sku:
                        data["products"][index_by_sku[key]] = product
                    else:
                        index_by_sku[key] = len(data["products"])
                        data["products"].append(product)
            return {"imported": len(cleaned), "count": len(data["products"])}

        return self._update_master(mutate)

    def load_settings(self) -> dict:
        with self.lock:
            return self._read_settings()

    def save_settings(self, incoming: dict) -> dict:
        if not isinstance(incoming, dict):
            raise ValueError("Cài đặt không hợp lệ")
        with self.lock:
            settings = self._coerce_settings(incoming, self._read_settings())
            self._write_json(self.settings_path, settings)
            return settings

    def _update_master(self, mutator):
        with self.lock:
            data = self._read_master()
            result = mutator(data)
            self._write_json(self.master_path, {"products": data["products"]})
            return copy.deepcopy(result)

    def _read_master(self) -> dict:
        self.data_dir.mkdir(parents=True, exist_ok=True)
        if not self.master_path.exists():
            data = default_master()
            self._write_json(self.master_path, data)
            return data
        try:
            raw = json.loads(self.master_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise RuntimeError("Không đọc được data/master.json") from exc
        return self._normalize(raw)

    def _normalize(self, raw: dict) -> dict:
        if not isinstance(raw, dict):
            raise RuntimeError("data/master.json không đúng định dạng")
        products = []
        for item in raw.get("products", []):
            if not isinstance(item, dict):
                continue
            payload = {
                "sku": item.get("sku", ""),
                "name": item.get("name", ""),
                "qr": item.get("qr") or item.get("qr_code") or "",
            }
            try:
                products.append(clean_product(payload))
            except ValueError as exc:
                sku = item.get("sku", "?")
                raise RuntimeError(f"Sản phẩm {sku} trong master data không hợp lệ: {exc}") from exc
        return {"products": products}

    def _read_settings(self) -> dict:
        settings = dict(DEFAULT_SETTINGS)
        if not self.settings_path.exists():
            return settings
        try:
            stored = json.loads(self.settings_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise RuntimeError("Không đọc được data/settings.json") from exc
        if isinstance(stored, dict):
            for key in DEFAULT_SETTINGS:
                if key in stored:
                    settings[key] = stored[key]
        return self._coerce_settings(settings, dict(DEFAULT_SETTINGS))

    def _coerce_settings(self, incoming: dict, base: dict) -> dict:
        settings = dict(base)
        if "printer" in incoming:
            settings["printer"] = str(incoming["printer"])[:200]
        settings["preset"] = "35x22"
        settings["label_width_mm"] = LABEL_WIDTH_MM
        settings["label_height_mm"] = LABEL_HEIGHT_MM
        try:
            if "copies" in incoming:
                settings["copies"] = int(incoming["copies"])
        except (TypeError, ValueError) as exc:
            raise ValueError("Cài đặt không hợp lệ") from exc
        settings["copies"] = normalize_copies(settings["copies"])
        return settings

    def _write_json(self, path: Path, payload: dict) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
